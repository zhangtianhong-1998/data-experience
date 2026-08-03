from __future__ import annotations

import copy
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from .utils import (
    clean_sql,
    file_record,
    make_ref,
    now_iso,
    read_json,
    short_hash,
    split_path,
    text_or_none,
)


EXECUTION_HEADERS = ["报表路径", "组件名称", "数据集路径", "执行SQL"]
LINEAGE_HEADERS = [
    "报表路径",
    "依赖的数据集路径",
    "绑定的列名",
    "依赖的物理表名",
    "依赖的数据库字段",
]
SYSTEM_COMPONENT = re.compile(r"^(?:图表|组件|筛选器)\s*\d*$", re.IGNORECASE)
DATE_PREFIX = re.compile(r"^[【\[]\d{6,8}[】\]]\s*")


def read_excel_rows(path: Path, headers: list[str]) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("缺少 openpyxl，请先安装 requirements.txt") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        actual = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as exc:
        raise ValueError(f"工作簿为空：{path}") from exc
    missing = [header for header in headers if header not in actual]
    if missing:
        raise ValueError(f"{path.name} 缺少表头：{missing}；实际表头：{actual}")
    positions = {header: actual.index(header) for header in headers}
    result: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows, start=2):
        record = {
            header: values[position] if position < len(values) else None
            for header, position in positions.items()
        }
        if any(text_or_none(value) is not None for value in record.values()):
            record["_row_number"] = row_number
            result.append(record)
    workbook.close()
    return result


def metadata_status(payload: Any) -> str:
    if not isinstance(payload, dict):
        return "结构异常：物理表元数据不是对象"
    issues: list[str] = []
    if text_or_none(payload.get("cn_name")) is None:
        issues.append("存在缺失值：表中文名 cn_name 缺失")
    fields = payload.get("fields")
    if fields is None:
        issues.append("存在缺失值：fields 缺失")
    elif not isinstance(fields, list):
        issues.append("结构异常：fields 不是数组")
    elif not fields:
        issues.append("存在缺失值：fields 为空")
    else:
        for index, field in enumerate(fields, start=1):
            if not isinstance(field, dict):
                issues.append(f"结构异常：第{index}个字段不是对象")
                continue
            name = text_or_none(field.get("field"))
            if name is None:
                issues.append(f"存在缺失值：第{index}个字段的 field 缺失")
            if text_or_none(field.get("comment")) is None:
                issues.append(f"存在缺失值：字段 {name or index} 的 comment 缺失")
    return "正常，未发现缺失值" if not issues else "；".join(issues)


def enrich_physical_metadata(source: dict[str, Any]) -> dict[str, Any]:
    enriched: dict[str, Any] = {}
    for table_name, raw_payload in source.items():
        if isinstance(raw_payload, dict):
            payload = copy.deepcopy(raw_payload)
            payload.pop("status", None)
            payload["status"] = metadata_status(payload)
        else:
            payload = {
                "raw_value": copy.deepcopy(raw_payload),
                "status": metadata_status(raw_payload),
            }
        enriched[str(table_name)] = payload
    return enriched


def _mapping_for_report(
    segments: list[str], mapping: dict[str, Any]
) -> tuple[str | None, dict[str, Any] | None]:
    if not segments:
        return None, None
    title = segments[-1]
    candidates = [title, DATE_PREFIX.sub("", title)]
    for candidate in candidates:
        if candidate in mapping:
            return candidate, copy.deepcopy(mapping[candidate])
    normalized = re.sub(r"[\s【】\[\]]+", "", title).lower()
    matches = [
        key
        for key in mapping
        if re.sub(r"[\s【】\[\]]+", "", str(key)).lower() == normalized
    ]
    if len(matches) == 1:
        return matches[0], copy.deepcopy(mapping[matches[0]])
    return None, None


def _scope_index(scope: dict[str, Any]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for department in scope.get("departments", []):
        for raw_path in department.get("report_paths", []):
            canonical, _, _ = split_path(raw_path)
            if canonical:
                result[canonical] = department
    return result


def _table_indexes(metadata: dict[str, Any]) -> tuple[dict[str, str], dict[str, list[str]]]:
    exact: dict[str, str] = {}
    leaf: dict[str, list[str]] = defaultdict(list)
    for name in metadata:
        key = str(name).strip().lower()
        exact[key] = name
        leaf[key.rsplit(".", 1)[-1]].append(name)
    return exact, leaf


def resolve_table_name(
    value: Any, exact: dict[str, str], leaf: dict[str, list[str]]
) -> tuple[str | None, str]:
    name = text_or_none(value)
    if name is None:
        return None, "missing_table"
    key = name.replace('"', "").replace("`", "").strip().lower()
    if key in exact:
        return exact[key], "matched_exact"
    matches = leaf.get(key.rsplit(".", 1)[-1], [])
    if len(matches) == 1:
        return matches[0], "matched_unique_leaf"
    if len(matches) > 1:
        return name, "ambiguous_leaf"
    return name, "metadata_unmatched"


def field_exists(metadata: dict[str, Any], table_name: str | None, field_name: str | None) -> bool:
    if not table_name or not field_name:
        return False
    payload = metadata.get(table_name)
    if not isinstance(payload, dict):
        return False
    return any(
        isinstance(field, dict) and text_or_none(field.get("field")) == field_name
        for field in payload.get("fields") or []
    )


def build_asset_graph(
    execution_excel: Path,
    table_metadata_path: Path,
    report_mapping_path: Path,
    lineage_excel: Path,
    scope_path: Path,
) -> tuple[dict[str, Any], dict[str, Any]]:
    execution_rows = read_excel_rows(execution_excel, EXECUTION_HEADERS)
    lineage_rows = read_excel_rows(lineage_excel, LINEAGE_HEADERS)
    source_metadata = read_json(table_metadata_path)
    report_mapping = read_json(report_mapping_path)
    scope = read_json(scope_path)
    metadata = enrich_physical_metadata(source_metadata)
    scope_by_report = _scope_index(scope)
    issues: list[dict[str, Any]] = []

    normalized_rows: list[dict[str, Any]] = []
    for row in execution_rows:
        row_number = row["_row_number"]
        report_path, report_segments, report_empty = split_path(row.get("报表路径"))
        dataset_path, dataset_segments, dataset_empty = split_path(row.get("数据集路径"))
        if report_path is None:
            report_key = f"__missing_report_row_{row_number}"
            issues.append(
                {
                    "source_ref": f"{execution_excel.name}#row={row_number}",
                    "code": "missing_report_path",
                    "message": "报表路径缺失",
                }
            )
        else:
            report_key = report_path
        if report_empty:
            issues.append(
                {
                    "source_ref": f"{execution_excel.name}#row={row_number}",
                    "code": "report_path_empty_segment",
                    "message": f"报表路径存在空片段：{report_empty}",
                }
            )
        if dataset_empty:
            issues.append(
                {
                    "source_ref": f"{execution_excel.name}#row={row_number}",
                    "code": "dataset_path_empty_segment",
                    "message": f"数据集路径存在空片段：{dataset_empty}",
                }
            )
        normalized_rows.append(
            {
                "source_ref": f"{execution_excel.name}#row={row_number}",
                "row_number": row_number,
                "report_key": report_key,
                "report_path_raw": row.get("报表路径"),
                "report_path": report_path,
                "report_segments": report_segments,
                "component_name": text_or_none(row.get("组件名称")),
                "dataset_path_raw": row.get("数据集路径"),
                "dataset_path": dataset_path,
                "dataset_segments": dataset_segments,
                "sql_raw": row.get("执行SQL"),
                "sql": clean_sql(row.get("执行SQL")),
            }
        )

    rows_by_report: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in normalized_rows:
        rows_by_report[row["report_key"]].append(row)

    reports: dict[str, dict[str, Any]] = {}
    query_lookup: dict[str, dict[str, Any]] = {}
    for report_key in sorted(rows_by_report):
        rows = rows_by_report[report_key]
        representative = next((row for row in rows if row["report_path"]), rows[0])
        canonical_path = representative["report_path"]
        report_ref = make_ref(
            "report", canonical_path or f"missing/{representative['row_number']}"
        )
        raw_variants = sorted(
            {
                str(row["report_path_raw"])
                for row in rows
                if row.get("report_path_raw") is not None
            }
        )
        mapping_key, mapping_payload = _mapping_for_report(
            representative["report_segments"], report_mapping
        )
        department = scope_by_report.get(canonical_path or "")
        report_issues: list[str] = []
        if canonical_path is None:
            report_issues.append("报表路径缺失，身份为临时引用")
        if mapping_payload is None:
            report_issues.append("未匹配报表使用的数据映射")
        if department is None:
            report_issues.append("未在实验范围中配置部门归属")

        component_groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        name_to_datasets: dict[str, set[str]] = defaultdict(set)
        for row in rows:
            component_key = row["component_name"] or "__missing_component_name__"
            dataset_key = row["dataset_path"] or "__missing_dataset_path__"
            component_groups[(component_key, dataset_key)].append(row)
            name_to_datasets[component_key].add(dataset_key)
        for name, datasets in name_to_datasets.items():
            if len(datasets) > 1:
                report_issues.append(
                    f"同名组件 {name} 关联多个数据集，已拆成独立组件：{sorted(datasets)}"
                )

        components: list[dict[str, Any]] = []
        sorted_groups = sorted(component_groups.items(), key=lambda item: item[0])
        for component_no, ((component_key, dataset_key), group_rows) in enumerate(
            sorted_groups, start=1
        ):
            component_ref = make_ref("component", report_ref.removeprefix("report://"), f"C{component_no:02d}")
            component_name = (
                None if component_key == "__missing_component_name__" else component_key
            )
            if component_name is None:
                name_kind = "missing"
            elif SYSTEM_COMPONENT.fullmatch(component_name):
                name_kind = "system_generated"
            else:
                name_kind = "business_named"
            dataset_path = None if dataset_key == "__missing_dataset_path__" else dataset_key
            dataset_ref = make_ref(
                "dataset",
                dataset_path
                or f"missing/{report_ref.removeprefix('report://')}/C{component_no:02d}",
            )
            dataset_segments = next(
                (row["dataset_segments"] for row in group_rows if row["dataset_path"]),
                [],
            )
            sql_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
            missing_sql_sources: list[str] = []
            for row in group_rows:
                if row["sql"]:
                    sql_groups[row["sql"]].append(row)
                else:
                    missing_sql_sources.append(row["source_ref"])
            queries: list[dict[str, Any]] = []
            for query_no, query_sql in enumerate(sorted(sql_groups), start=1):
                query_rows = sql_groups[query_sql]
                query_ref = make_ref(
                    "query",
                    component_ref.removeprefix("component://"),
                    f"Q{query_no:02d}",
                )
                query = {
                    "query_ref": query_ref,
                    "query_no": query_no,
                    "sql": query_sql,
                    "occurrence_count": len(query_rows),
                    "source_refs": sorted(row["source_ref"] for row in query_rows),
                    "fingerprint": short_hash(query_sql, 24),
                    "status": "待SQL解析",
                }
                queries.append(query)
                query_lookup[query_ref] = query
            dataset_issues: list[str] = []
            if dataset_path is None:
                dataset_issues.append("数据集路径缺失")
            if missing_sql_sources:
                dataset_issues.append(f"存在{len(missing_sql_sources)}条SQL缺失记录")
            components.append(
                {
                    "component_ref": component_ref,
                    "component_no": component_no,
                    "component_name": component_name,
                    "name_kind": name_kind,
                    "source_refs": sorted(row["source_ref"] for row in group_rows),
                    "dataset": {
                        "dataset_ref": dataset_ref,
                        "dataset_path": dataset_path,
                        "dataset_path_segments": dataset_segments,
                        "identity_status": "provisional_path" if dataset_path else "provisional_missing",
                        "queries": queries,
                        "missing_sql_source_refs": missing_sql_sources,
                        "status": "正常" if not dataset_issues else "；".join(dataset_issues),
                    },
                    "status": "正常" if component_name else "组件名称缺失",
                }
            )

        reports[report_ref] = {
            "report_ref": report_ref,
            "source_id": mapping_payload.get("table_id") if mapping_payload else None,
            "identity_status": "source_mapped" if mapping_payload and mapping_payload.get("table_id") else "provisional_path",
            "report_path": canonical_path,
            "report_path_segments": representative["report_segments"],
            "report_path_raw_variants": raw_variants,
            "mapping": {
                "mapping_key": mapping_key,
                **(mapping_payload or {}),
            },
            "department_ref": department.get("department_ref") if department else None,
            "department_name": department.get("name") if department else None,
            "department_assignment_method": department.get("assignment_method") if department else None,
            "components": components,
            "status": "正常" if not report_issues else "；".join(report_issues),
        }

    exact_tables, leaf_tables = _table_indexes(metadata)
    lineage_records: list[dict[str, Any]] = []
    report_by_path = {
        report["report_path"]: report_ref
        for report_ref, report in reports.items()
        if report.get("report_path")
    }
    for row in lineage_rows:
        row_number = row["_row_number"]
        report_path, _, _ = split_path(row.get("报表路径"))
        dataset_path, _, _ = split_path(row.get("依赖的数据集路径"))
        table_name, table_match = resolve_table_name(
            row.get("依赖的物理表名"), exact_tables, leaf_tables
        )
        field_name = text_or_none(row.get("依赖的数据库字段"))
        report_ref = report_by_path.get(report_path or "")
        statuses: list[str] = []
        if report_ref is None:
            statuses.append("报表未匹配")
        if table_match not in {"matched_exact", "matched_unique_leaf"}:
            statuses.append(f"物理表匹配状态：{table_match}")
        if field_name is None:
            statuses.append("数据库字段缺失")
        elif not field_exists(metadata, table_name, field_name):
            statuses.append("数据库字段不在物理元数据中")
        table_ref = make_ref("table", table_name) if table_name else None
        field_ref = (
            make_ref("field", table_name, field_name)
            if table_name and field_name
            else None
        )
        lineage_records.append(
            {
                "lineage_ref": make_ref("lineage", f"row-{row_number}"),
                "source_ref": f"{lineage_excel.name}#row={row_number}",
                "report_ref": report_ref,
                "report_path": report_path,
                "dataset_path": dataset_path,
                "bound_column_name": text_or_none(row.get("绑定的列名")),
                "table_name": table_name,
                "table_ref": table_ref,
                "field_name": field_name,
                "field_ref": field_ref,
                "status": "正常" if not statuses else "；".join(statuses),
            }
        )

    selected_department_ref = scope.get("selected_department_ref")
    asset_graph = {
        "schema_version": "asset-graph-v1",
        "generated_at": now_iso(),
        "scope": {
            "group": scope.get("group"),
            "selected_department_ref": selected_department_ref,
            "policy": scope.get("policy", {}),
        },
        "physical_metadata": metadata,
        "reports": reports,
        "field_lineage": lineage_records,
        "issues": issues,
    }
    snapshot = {
        "schema_version": "source-snapshot-v1",
        "generated_at": now_iso(),
        "sources": [
            file_record(execution_excel, len(execution_rows)),
            file_record(table_metadata_path, len(source_metadata)),
            file_record(report_mapping_path, len(report_mapping)),
            file_record(lineage_excel, len(lineage_rows)),
            file_record(scope_path, len(scope.get("departments", []))),
        ],
        "input_issue_count": len(issues),
        "issues": issues,
    }
    return snapshot, asset_graph

